"""Lookup and reference Excel functions."""

from __future__ import annotations

import math
from typing import Any, List, Optional

from simulation.ast_nodes import ExcelError, NA_ERROR, VALUE_ERROR, REF_ERROR
from simulation.functions import register


def _to_2d(val: Any) -> List[List[Any]]:
    """Ensure val is a 2D list (list of rows)."""
    if not isinstance(val, list):
        return [[val]]
    if not val:
        return [[]]
    if isinstance(val[0], list):
        return val
    return [val]


def _match_value(lookup_val: Any, cell_val: Any, match_type: int) -> bool:
    """Compare two values for MATCH/VLOOKUP purposes."""
    if match_type == 0:
        if isinstance(lookup_val, str) and isinstance(cell_val, str):
            return lookup_val.upper() == cell_val.upper()
        return lookup_val == cell_val
    return False


@register("VLOOKUP")
def _vlookup(evaluator, args):
    lookup_val = evaluator.resolve_arg(args[0])
    table = _to_2d(evaluator.resolve_arg(args[1]))
    col_index = int(evaluator.resolve_arg(args[2]))
    range_lookup = True
    if len(args) > 3:
        rl = evaluator.resolve_arg(args[3])
        if isinstance(rl, bool):
            range_lookup = rl
        elif isinstance(rl, (int, float)):
            range_lookup = rl != 0

    if col_index < 1 or not table or col_index > len(table[0]):
        return REF_ERROR

    if range_lookup:
        # Approximate match: table must be sorted ascending by first column
        best_row = None
        for row in table:
            if not row:
                continue
            cell = row[0]
            if cell is None:
                continue
            try:
                if isinstance(lookup_val, (int, float)) and isinstance(cell, (int, float)):
                    if cell <= lookup_val:
                        best_row = row
                    else:
                        break
                elif isinstance(lookup_val, str) and isinstance(cell, str):
                    if cell.upper() <= lookup_val.upper():
                        best_row = row
                    else:
                        break
            except TypeError:
                continue
        if best_row is None:
            return NA_ERROR
        return best_row[col_index - 1] if col_index - 1 < len(best_row) else NA_ERROR
    else:
        # Exact match
        for row in table:
            if not row:
                continue
            if _match_value(lookup_val, row[0], 0):
                return row[col_index - 1] if col_index - 1 < len(row) else NA_ERROR
        return NA_ERROR


@register("HLOOKUP")
def _hlookup(evaluator, args):
    lookup_val = evaluator.resolve_arg(args[0])
    table = _to_2d(evaluator.resolve_arg(args[1]))
    row_index = int(evaluator.resolve_arg(args[2]))
    range_lookup = True
    if len(args) > 3:
        rl = evaluator.resolve_arg(args[3])
        if isinstance(rl, bool):
            range_lookup = rl
        elif isinstance(rl, (int, float)):
            range_lookup = rl != 0

    if row_index < 1 or not table or row_index > len(table):
        return REF_ERROR

    first_row = table[0] if table else []

    if range_lookup:
        best_col = None
        for ci, cell in enumerate(first_row):
            if cell is None:
                continue
            try:
                if isinstance(lookup_val, (int, float)) and isinstance(cell, (int, float)):
                    if cell <= lookup_val:
                        best_col = ci
                    else:
                        break
                elif isinstance(lookup_val, str) and isinstance(cell, str):
                    if cell.upper() <= lookup_val.upper():
                        best_col = ci
                    else:
                        break
            except TypeError:
                continue
        if best_col is None:
            return NA_ERROR
        target_row = table[row_index - 1]
        return target_row[best_col] if best_col < len(target_row) else NA_ERROR
    else:
        for ci, cell in enumerate(first_row):
            if _match_value(lookup_val, cell, 0):
                target_row = table[row_index - 1]
                return target_row[ci] if ci < len(target_row) else NA_ERROR
        return NA_ERROR


@register("INDEX")
def _index(evaluator, args):
    array = _to_2d(evaluator.resolve_arg(args[0]))
    row_num = evaluator.resolve_arg(args[1]) if len(args) > 1 else None
    col_num = evaluator.resolve_arg(args[2]) if len(args) > 2 else None

    if row_num is not None:
        row_num = int(row_num)
    if col_num is not None:
        col_num = int(col_num)

    if not array:
        return REF_ERROR

    # Row 0 means entire column, col 0 means entire row
    if row_num == 0 and col_num is not None and col_num > 0:
        result = []
        for r in array:
            if col_num - 1 < len(r):
                result.append(r[col_num - 1])
            else:
                result.append(None)
        return result

    if col_num == 0 and row_num is not None and row_num > 0:
        if row_num - 1 < len(array):
            return array[row_num - 1]
        return REF_ERROR

    if row_num is None or row_num == 0:
        row_num = 1
    if col_num is None or col_num == 0:
        col_num = 1

    if row_num < 1 or row_num > len(array):
        return REF_ERROR
    row = array[row_num - 1]
    if col_num < 1 or col_num > len(row):
        return REF_ERROR
    return row[col_num - 1]


@register("MATCH")
def _match(evaluator, args):
    lookup_val = evaluator.resolve_arg(args[0])
    array_raw = evaluator.resolve_arg(args[1])
    match_type = int(evaluator.resolve_arg(args[2])) if len(args) > 2 else 1

    # Flatten to 1D
    if isinstance(array_raw, list):
        flat = []
        for item in array_raw:
            if isinstance(item, list):
                flat.extend(item)
            else:
                flat.append(item)
    else:
        flat = [array_raw]

    if match_type == 0:
        for i, cell in enumerate(flat):
            if _match_value(lookup_val, cell, 0):
                return float(i + 1)
        return NA_ERROR

    elif match_type == 1:
        # Largest value <= lookup (array sorted ascending)
        best = None
        for i, cell in enumerate(flat):
            if cell is None:
                continue
            try:
                if isinstance(lookup_val, (int, float)) and isinstance(cell, (int, float)):
                    if cell <= lookup_val:
                        best = i
                    else:
                        break
                elif isinstance(lookup_val, str) and isinstance(cell, str):
                    if cell.upper() <= lookup_val.upper():
                        best = i
                    else:
                        break
            except TypeError:
                continue
        return float(best + 1) if best is not None else NA_ERROR

    elif match_type == -1:
        # Smallest value >= lookup (array sorted descending)
        best = None
        for i, cell in enumerate(flat):
            if cell is None:
                continue
            try:
                if isinstance(lookup_val, (int, float)) and isinstance(cell, (int, float)):
                    if cell >= lookup_val:
                        best = i
                    else:
                        break
                elif isinstance(lookup_val, str) and isinstance(cell, str):
                    if cell.upper() >= lookup_val.upper():
                        best = i
                    else:
                        break
            except TypeError:
                continue
        return float(best + 1) if best is not None else NA_ERROR

    return NA_ERROR


@register("CHOOSE")
def _choose(evaluator, args):
    idx = int(evaluator.resolve_arg(args[0]))
    if idx < 1 or idx >= len(args):
        return VALUE_ERROR
    return evaluator.resolve_arg(args[idx])


@register("ROW")
def _row(evaluator, args):
    if not args:
        return float(evaluator.current_row)
    val = args[0]
    from simulation.ast_nodes import CellRef, RangeRef
    if isinstance(val, CellRef):
        return float(val.row)
    if isinstance(val, RangeRef):
        return float(val.start.row)
    return float(evaluator.current_row)


@register("COLUMN")
def _column(evaluator, args):
    if not args:
        return float(evaluator.current_col_index)
    val = args[0]
    from simulation.ast_nodes import CellRef, RangeRef
    from openpyxl.utils import column_index_from_string
    if isinstance(val, CellRef):
        return float(column_index_from_string(val.col))
    if isinstance(val, RangeRef):
        return float(column_index_from_string(val.start.col))
    return float(evaluator.current_col_index)


@register("ROWS")
def _rows(evaluator, args):
    val = evaluator.resolve_arg(args[0])
    data = _to_2d(val)
    return float(len(data))


@register("COLUMNS")
def _columns(evaluator, args):
    val = evaluator.resolve_arg(args[0])
    data = _to_2d(val)
    return float(len(data[0])) if data and data[0] else 0.0


@register("OFFSET")
def _offset(evaluator, args):
    """OFFSET(reference, rows, cols, [height], [width])

    Returns a range shifted from the reference. We resolve the base cell
    and shift by the given row/col offsets, returning a single value or
    a 2D array for height/width > 1.
    """
    base = args[0]  # raw AST node
    row_off = int(evaluator.resolve_arg(args[1]))
    col_off = int(evaluator.resolve_arg(args[2]))
    height = int(evaluator.resolve_arg(args[3])) if len(args) > 3 else 1
    width = int(evaluator.resolve_arg(args[4])) if len(args) > 4 else 1

    from simulation.ast_nodes import CellRef
    from openpyxl.utils import column_index_from_string, get_column_letter

    if not isinstance(base, CellRef):
        return REF_ERROR

    sheet = base.sheet or evaluator.current_sheet
    base_col_idx = column_index_from_string(base.col)
    new_row = base.row + row_off
    new_col_idx = base_col_idx + col_off

    if new_row < 1 or new_col_idx < 1:
        return REF_ERROR

    if height == 1 and width == 1:
        new_col = get_column_letter(new_col_idx)
        return evaluator.get_cell_value(sheet, new_col, new_row)

    result = []
    for r in range(height):
        row_data = []
        for c in range(width):
            col_letter = get_column_letter(new_col_idx + c)
            row_data.append(evaluator.get_cell_value(sheet, col_letter, new_row + r))
        result.append(row_data)
    return result
