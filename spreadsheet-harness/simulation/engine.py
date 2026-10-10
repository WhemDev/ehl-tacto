"""Core Excel simulation engine with dependency-graph-based cascade recalculation.

This module replaces the ``formulas`` library.  It loads an Excel workbook
via openpyxl, parses every formula into an AST, builds a dependency graph,
and supports *what-if* analysis: change input cells, recalculate only the
affected downstream cells in topological order, then optionally restore
the original state.
"""

from __future__ import annotations

import copy
import os
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from simulation.ast_nodes import (
    ASTNode,
    BinOp,
    BoolLiteral,
    CellRef,
    ErrorValue,
    ExcelError,
    FuncCall,
    NumberLiteral,
    Percent,
    RangeRef,
    StringLiteral,
    UnaryOp,
    VALUE_ERROR,
    DIV0_ERROR,
    REF_ERROR,
    NAME_ERROR,
)
from simulation.dependency_graph import CellKey, DependencyGraph, make_key
from simulation.parser import parse_formula
from simulation.functions import FUNCTION_REGISTRY, ensure_loaded
# The engine's own date<->serial helper (date_funcs.py:21-26), reused rather
# than reimplemented here — it already encodes the 1899-12-30 epoch and the
# "1900 is a leap year" Lotus 1-2-3 bug exactly as the TS port's
# dateToSerial (lib/formula/functions/date.ts) does; a second implementation
# would risk silently disagreeing with the first. No circular import: this
# module (simulation.functions.date_funcs) only imports from
# simulation.ast_nodes and simulation.functions (the registry), neither of
# which imports simulation.engine.
from simulation.functions.date_funcs import _date_to_serial


class ExcelSimulator:
    """Pure-Python Excel formula engine with cascade recalculation.

    Usage::

        sim = ExcelSimulator("Excaliba.xlsx")
        results = sim.evaluate_with_changes(
            changes={"UrunPlan": {"B6": 107142}},
            targets={"UrunPlan": ["B7", "B9"]},
        )
    """

    def __init__(self, file_path: str):
        self.file_path = file_path
        # {CellKey: parsed ASTNode}
        self._formulas: Dict[CellKey, ASTNode] = {}
        # {CellKey: current value (scalar)}
        self._values: Dict[CellKey, Any] = {}
        # All sheet names (upper-cased for lookup, original for display)
        self._sheet_names: Dict[str, str] = {}  # UPPER -> original
        self._graph = DependencyGraph()
        self._loaded = False
        # Evaluation context (set during evaluation for function callbacks)
        self._eval_sheet: str = ""
        self._eval_col: str = ""
        self._eval_row: int = 0

        ensure_loaded()

    # ── Loading ──────────────────────────────────────────────

    def _load(self) -> None:
        """Load workbook, parse formulas, build dependency graph."""
        if self._loaded:
            return

        wb = openpyxl.load_workbook(self.file_path, data_only=False)
        wb_data = openpyxl.load_workbook(self.file_path, data_only=True)

        try:
            for ws in wb.worksheets:
                sheet_upper = ws.title.upper()
                self._sheet_names[sheet_upper] = ws.title
                ws_data = wb_data[ws.title]

                for row in ws.iter_rows(
                    min_row=1, max_row=ws.max_row,
                    min_col=1, max_col=ws.max_column,
                    values_only=False,
                ):
                    for cell in row:
                        col = get_column_letter(cell.column)
                        key = make_key(ws.title, col, cell.row)
                        raw = cell.value

                        if isinstance(raw, str) and raw.startswith("="):
                            try:
                                ast = parse_formula(raw)
                                self._formulas[key] = ast
                                refs = self._extract_refs(ast, sheet_upper)
                                self._graph.add_formula_cell(key, refs)
                            except Exception:
                                # Unparseable formula — store cached value
                                cached = ws_data.cell(
                                    row=cell.row, column=cell.column
                                ).value
                                self._values[key] = cached
                        else:
                            self._values[key] = raw

        finally:
            wb.close()
            wb_data.close()

        self._loaded = True
        self._full_recalculate()
        print(f"    [ExcelSimulator] Loaded: "
              f"{len(self._formulas)} formulas, "
              f"{len(self._values)} values, "
              f"{self._graph.stats()['total_edges']} edges")

    def _full_recalculate(self) -> None:
        """Evaluate all formula cells in dependency order."""
        self._graph.mark_all_dirty()
        self._recalculate()

    # ── Reference extraction (AST -> CellKeys) ──────────────

    def _extract_refs(self, node: ASTNode, default_sheet: str) -> Set[CellKey]:
        """Walk AST and collect all cell references as CellKeys."""
        refs: Set[CellKey] = set()
        self._walk_refs(node, default_sheet, refs)
        return refs

    def _walk_refs(self, node: ASTNode, default_sheet: str, refs: Set[CellKey]) -> None:
        if isinstance(node, CellRef):
            sheet = (node.sheet or default_sheet).upper()
            refs.add(make_key(sheet, node.col, node.row))

        elif isinstance(node, RangeRef):
            sheet = (node.start.sheet or default_sheet).upper()
            start_col = column_index_from_string(node.start.col)
            end_col = column_index_from_string(node.end.col)
            for c in range(start_col, end_col + 1):
                for r in range(node.start.row, node.end.row + 1):
                    refs.add(make_key(sheet, get_column_letter(c), r))

        elif isinstance(node, BinOp):
            self._walk_refs(node.left, default_sheet, refs)
            self._walk_refs(node.right, default_sheet, refs)

        elif isinstance(node, UnaryOp):
            self._walk_refs(node.operand, default_sheet, refs)

        elif isinstance(node, Percent):
            self._walk_refs(node.operand, default_sheet, refs)

        elif isinstance(node, FuncCall):
            for arg in node.args:
                self._walk_refs(arg, default_sheet, refs)

    # ── Evaluation ───────────────────────────────────────────

    def _recalculate(self) -> None:
        """Recalculate all dirty cells in topological order."""
        order = self._graph.topological_order()
        for key in order:
            if key in self._formulas:
                self._eval_sheet = key[0]
                self._eval_col = key[1]
                self._eval_row = key[2]
                try:
                    val = self._eval_node(self._formulas[key], key[0])
                except Exception:
                    val = VALUE_ERROR
                self._values[key] = val
            self._graph.clear_dirty(key)

    def _eval_node(self, node: ASTNode, sheet: str) -> Any:
        """Recursively evaluate an AST node."""
        if isinstance(node, NumberLiteral):
            return node.value

        if isinstance(node, StringLiteral):
            return node.value

        if isinstance(node, BoolLiteral):
            return node.value

        if isinstance(node, ErrorValue):
            return ExcelError(node.error_type)

        if isinstance(node, CellRef):
            ref_sheet = (node.sheet or sheet).upper()
            key = make_key(ref_sheet, node.col, node.row)
            return self._values.get(key)

        if isinstance(node, RangeRef):
            return self._resolve_range(node, sheet)

        if isinstance(node, BinOp):
            return self._eval_binop(node, sheet)

        if isinstance(node, UnaryOp):
            val = self._eval_node(node.operand, sheet)
            if isinstance(val, ExcelError):
                return val
            if node.op == "-":
                return -self._to_number(val)
            return val

        if isinstance(node, Percent):
            val = self._eval_node(node.operand, sheet)
            if isinstance(val, ExcelError):
                return val
            return self._to_number(val) / 100.0

        if isinstance(node, FuncCall):
            return self._eval_func(node, sheet)

        return VALUE_ERROR

    def _eval_binop(self, node: BinOp, sheet: str) -> Any:
        left = self._eval_node(node.left, sheet)
        right = self._eval_node(node.right, sheet)

        if isinstance(left, ExcelError):
            return left
        if isinstance(right, ExcelError):
            return right

        op = node.op

        if op == "&":
            return self._to_str(left) + self._to_str(right)

        # Comparison operators
        if op in ("=", "<>", "<", ">", "<=", ">="):
            return self._compare(left, right, op)

        # Arithmetic
        l = self._to_number(left)
        r = self._to_number(right)
        if isinstance(l, ExcelError):
            return l
        if isinstance(r, ExcelError):
            return r

        if op == "+":
            return l + r
        if op == "-":
            return l - r
        if op == "*":
            return l * r
        if op == "/":
            if r == 0:
                return DIV0_ERROR
            return l / r
        if op == "^":
            try:
                return l ** r
            except (OverflowError, ValueError):
                return ExcelError("#NUM!")

        return VALUE_ERROR

    def _compare(self, left: Any, right: Any, op: str) -> bool:
        # Excel comparison: numbers vs numbers, strings vs strings (case-insensitive)
        if isinstance(left, (int, float)) and not isinstance(left, bool):
            if isinstance(right, (int, float)) and not isinstance(right, bool):
                l, r = float(left), float(right)
            elif right is None:
                l, r = float(left), 0.0
            else:
                # number < text in Excel
                if op == "<":
                    return True
                if op == ">":
                    return False
                if op == "=":
                    return False
                if op == "<>":
                    return True
                if op == "<=":
                    return True
                if op == ">=":
                    return False
                return False
        elif isinstance(left, str) and isinstance(right, str):
            l, r = left.upper(), right.upper()
        elif left is None and isinstance(right, (int, float)):
            l, r = 0.0, float(right)
        elif left is None and right is None:
            l, r = 0, 0
        elif isinstance(left, bool) or isinstance(right, bool):
            l = 1 if left else 0
            r = 1 if right else 0
        else:
            l, r = str(left) if left else "", str(right) if right else ""

        if op == "=":
            return l == r
        if op == "<>":
            return l != r
        if op == "<":
            return l < r
        if op == ">":
            return l > r
        if op == "<=":
            return l <= r
        if op == ">=":
            return l >= r
        return False

    def _resolve_range(self, node: RangeRef, sheet: str) -> List[List[Any]]:
        """Resolve a range reference to a 2D list of values."""
        ref_sheet = (node.start.sheet or sheet).upper()
        start_col = column_index_from_string(node.start.col)
        end_col = column_index_from_string(node.end.col)
        start_row = node.start.row
        end_row = node.end.row

        result = []
        for r in range(start_row, end_row + 1):
            row_data = []
            for c in range(start_col, end_col + 1):
                key = make_key(ref_sheet, get_column_letter(c), r)
                row_data.append(self._values.get(key))
            result.append(row_data)
        return result

    def _eval_func(self, node: FuncCall, sheet: str) -> Any:
        """Evaluate a function call using the function registry."""
        func = FUNCTION_REGISTRY.get(node.name)
        if func is None:
            return NAME_ERROR

        evaluator = _FuncEvaluator(self, sheet)
        try:
            return func(evaluator, node.args)
        except Exception:
            return VALUE_ERROR

    @staticmethod
    def _to_number(val: Any) -> float | ExcelError:
        if isinstance(val, ExcelError):
            return val
        if val is None:
            return 0.0
        if isinstance(val, bool):
            return 1.0 if val else 0.0
        if isinstance(val, (int, float)):
            return float(val)
        if isinstance(val, str):
            try:
                return float(val)
            except ValueError:
                return VALUE_ERROR
        # datetime is a subclass of date, so it must be checked first —
        # _date_to_serial requires a plain date (subtracting it from another
        # date works; subtracting a date from a datetime raises TypeError).
        # Same isinstance ordering as _coerce_date in date_funcs.py.
        if isinstance(val, datetime):
            return _date_to_serial(val.date())
        if isinstance(val, date):
            return _date_to_serial(val)
        return VALUE_ERROR

    @staticmethod
    def _to_str(val: Any) -> str:
        if val is None:
            return ""
        if isinstance(val, bool):
            return "TRUE" if val else "FALSE"
        if isinstance(val, float) and val == int(val):
            return str(int(val))
        return str(val)

    # ── Public API ───────────────────────────────────────────

    def get_value(self, sheet: str, cell: str) -> Any:
        """Get the current computed value of a cell."""
        self._load()
        col, row = self._parse_cell_addr(cell)
        key = make_key(sheet, col, row)
        return self._values.get(key)

    def set_value(self, sheet: str, cell: str, value: Any) -> None:
        """Write a value to a cell and mark dependents dirty."""
        self._load()
        col, row = self._parse_cell_addr(cell)
        key = make_key(sheet, col, row)
        self._values[key] = value
        # If this was a formula cell, it's now a constant
        if key in self._formulas:
            del self._formulas[key]
            self._graph.remove_formula(key)
        self._graph.mark_dirty(key)

    def recalculate(self) -> None:
        """Recalculate all dirty cells."""
        self._load()
        self._recalculate()

    def evaluate_with_changes(
        self,
        changes: Dict[str, Dict[str, Any]],
        targets: Dict[str, List[str]],
    ) -> Dict[str, Dict[str, Any]]:
        """Apply changes, recalculate, read targets, then restore original state.

        This is the primary API for what-if analysis.  It is
        non-destructive: the simulator state is restored after reading
        target values.

        Parameters:
            changes: ``{"SheetName": {"CellAddr": value, ...}, ...}``
            targets: ``{"SheetName": ["CellAddr", ...], ...}``

        Returns:
            ``{"SheetName": {"CellAddr": computed_value, ...}, ...}``
        """
        self._load()

        # Snapshot current state
        snap = self._snapshot()

        try:
            # Apply all changes
            for sheet, cells in changes.items():
                for cell_addr, value in cells.items():
                    col, row = self._parse_cell_addr(cell_addr)
                    key = make_key(sheet, col, row)
                    self._values[key] = value
                    # Remove formula if overwriting
                    if key in self._formulas:
                        del self._formulas[key]
                        self._graph.remove_formula(key)
                    self._graph.mark_dirty(key)

            # Recalculate
            self._recalculate()

            # Read targets
            results: Dict[str, Dict[str, Any]] = {}
            for sheet, cell_list in targets.items():
                results[sheet] = {}
                for cell_addr in cell_list:
                    col, row = self._parse_cell_addr(cell_addr)
                    key = make_key(sheet, col, row)
                    val = self._values.get(key)
                    results[sheet][cell_addr] = self._extract_scalar(val)

            return results

        finally:
            # Restore original state
            self._restore(snap)

    def find_optimal_value(
        self,
        write_sheet: str,
        write_cell: str,
        candidates: List[Any],
        check_sheet: str,
        check_cell: str,
        min_val: float,
        max_val: float,
        other_changes: Optional[Dict[str, Dict[str, Any]]] = None,
    ) -> Optional[Any]:
        """Find the smallest candidate that puts check_cell in [min_val, max_val].

        Parameters:
            write_sheet/write_cell: Cell to write candidate values to.
            candidates: Values to try (should be sorted ascending).
            check_sheet/check_cell: Cell whose computed value is checked.
            min_val/max_val: Target range for check_cell.
            other_changes: Additional changes applied alongside each candidate.

        Returns:
            The first candidate whose check_cell result is in range, or None.
        """
        self._load()

        for candidate in candidates:
            changes: Dict[str, Dict[str, Any]] = {
                write_sheet: {write_cell: candidate}
            }
            if other_changes:
                for s, c in other_changes.items():
                    changes.setdefault(s, {}).update(c)

            results = self.evaluate_with_changes(
                changes=changes,
                targets={check_sheet: [check_cell]},
            )

            val = results.get(check_sheet, {}).get(check_cell)
            if val is not None and min_val <= val <= max_val:
                return candidate

        return None

    def get_cell_value(self, sheet: str, col: str, row: int) -> Any:
        """Direct cell value access for function implementations."""
        key = make_key(sheet, col, row)
        return self._values.get(key)

    def dump_all(self) -> Dict[str, Any]:
        """Every computed cell, JSON-safe, keyed "Sheet!A1".

        Read-only introspection for the TypeScript port's oracle test — never
        used by evaluation itself, so it cannot change what this engine
        computes, only how a result is read out.
        """
        self._load()
        out: Dict[str, Any] = {}
        for (sheet, col, row), val in self._values.items():
            display_sheet = self._sheet_names.get(sheet, sheet)
            addr = f"{display_sheet}!{col}{row}"
            if isinstance(val, ExcelError):
                out[addr] = f"__error__{val.error_type}"
            elif val is None:
                out[addr] = None
            elif isinstance(val, bool):
                out[addr] = val
            elif isinstance(val, (int, float)):
                out[addr] = float(val)
            else:
                out[addr] = str(val)
        return out

    @property
    def current_sheet(self) -> str:
        return self._eval_sheet

    @property
    def current_row(self) -> int:
        return self._eval_row

    @property
    def current_col_index(self) -> int:
        try:
            return column_index_from_string(self._eval_col)
        except Exception:
            return 1

    def reset(self) -> None:
        """Force a full reload on next access."""
        self._loaded = False
        self._formulas.clear()
        self._values.clear()
        self._sheet_names.clear()
        self._graph = DependencyGraph()

    # ── Snapshot / Restore ───────────────────────────────────

    def _snapshot(self) -> dict:
        return {
            "values": dict(self._values),
            "formulas": dict(self._formulas),
            "graph_formula_cells": set(self._graph.formula_cells),
        }

    def _restore(self, snap: dict) -> None:
        self._values = snap["values"]
        self._formulas = snap["formulas"]
        # Rebuild graph from formulas
        self._graph = DependencyGraph()
        for key, ast in self._formulas.items():
            refs = self._extract_refs(ast, key[0])
            self._graph.add_formula_cell(key, refs)
        self._graph.clear_all_dirty()

    # ── Helpers ──────────────────────────────────────────────

    @staticmethod
    def _parse_cell_addr(addr: str) -> Tuple[str, int]:
        """Parse 'B6' -> ('B', 6)."""
        addr = addr.upper().replace("$", "")
        i = 0
        while i < len(addr) and addr[i].isalpha():
            i += 1
        col = addr[:i]
        row = int(addr[i:])
        return col, row

    @staticmethod
    def _extract_scalar(val: Any) -> Optional[float]:
        """Convert a cell value to float for external consumers."""
        if val is None:
            return None
        if isinstance(val, ExcelError):
            return None
        if isinstance(val, bool):
            return float(val)
        if isinstance(val, (int, float)):
            return float(val)
        if isinstance(val, str):
            try:
                return float(val)
            except ValueError:
                return None
        return None


class _FuncEvaluator:
    """Adapter that function implementations use to resolve arguments.

    Functions receive ``(evaluator, args)`` where ``evaluator`` is one of
    these.  They can call ``evaluator.resolve_arg(ast_node)`` to evaluate
    a function argument, or access range data and cell values.
    """

    def __init__(self, sim: ExcelSimulator, sheet: str):
        self._sim = sim
        self._sheet = sheet

    def resolve_arg(self, node: ASTNode) -> Any:
        """Evaluate an AST node in the current sheet context.

        Range nodes return 2D lists; everything else returns scalars.
        """
        return self._sim._eval_node(node, self._sheet)

    def get_cell_value(self, sheet: str, col: str, row: int) -> Any:
        return self._sim.get_cell_value(sheet, col, row)

    @property
    def current_sheet(self) -> str:
        return self._sheet

    @property
    def current_row(self) -> int:
        return self._sim._eval_row

    @property
    def current_col_index(self) -> int:
        return self._sim.current_col_index
